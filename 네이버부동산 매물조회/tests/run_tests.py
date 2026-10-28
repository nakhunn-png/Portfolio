# -*- coding: utf-8 -*-
"""
오프라인 통합 테스트 (인터넷 불필요)
- tests/fixtures.json 에 저장된 '실제 네이버 응답'(2026-10-24 수집)을 HTTP 대신 돌려줘서
  크롤러(백엔드)와 화면(프론트엔드)을 끝까지 검증함
- 실행:  python tests/run_tests.py          (프로젝트 폴더에서)
- 화면 캡처는 tests/_out/ 에 저장됨

주의: fixtures에 없는 주소를 요청하면 TimeoutError가 나며 'missing'에 표시됨.
      새 기능 테스트에 새 응답이 필요하면 fixtures.json에 {경로: {status, text}} 형태로 추가.
"""
import json
import os
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "src")
OUT = os.path.join(HERE, "_out")
os.makedirs(OUT, exist_ok=True)
os.environ["APPDATA"] = os.path.join(OUT, "appdata")          # 장바구니 저장 위치를 테스트용으로 분리
shutil.rmtree(os.environ["APPDATA"], ignore_errors=True)
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")         # 창을 띄우지 않고 테스트
if sys.platform == "win32":
    os.environ.setdefault("QT_QPA_FONTDIR", r"C:\Windows\Fonts")  # offscreen에서 한글 폰트
sys.path.insert(0, SRC)

FX = json.load(open(os.path.join(HERE, "fixtures.json"), encoding="utf-8"))
MISSING = []


class FakeResp:
    def __init__(self, status, text):
        self.status_code, self.text = status, text

    def json(self):
        return json.loads(self.text)


def fake_get(self, url, headers=None, timeout=None, **kw):
    p = url.replace("https://new.land.naver.com", "")
    if p not in FX:
        MISSING.append(p)
        raise TimeoutError("fixture missing: " + p)
    return FakeResp(FX[p]["status"], FX[p]["text"])


try:
    from curl_cffi import requests as http_lib
except ImportError:
    import requests as http_lib
http_lib.Session.get = fake_get

import naver_land_crawler as c  # noqa: E402

c.random.uniform = lambda a, b: 0
c.time.sleep = lambda s: None
_orig_wait = c.NaverLand._wait
c.NaverLand._wait = lambda self, sec: _orig_wait(self, 0)   # 대기 없이 (중지 확인은 그대로)

RESULTS = []


def check(name, cond, detail=""):
    RESULTS.append((name, bool(cond)))
    print(("  [OK]   " if cond else "  [FAIL] ") + name + (f"  ({detail})" if detail else ""))


def test_backend():
    print("\n[백엔드]")
    nl = c.NaverLand()
    cx = nl.search_complexes("압구정 현대")
    check("단지 검색: 압구정 현대 8개", len(cx) == 8, f"{len(cx)}개")
    cx = nl.search_complexes("위례")
    groups = {}
    for x in cx:
        groups[c.property_group(x)] = groups.get(c.property_group(x), 0) + 1
    check("단지 검색: 위례 3페이지 110개 (아파트 86/오피스텔 24)",
          len(cx) == 110 and groups == {"아파트": 86, "오피스텔": 24}, str(groups))
    arts = nl.get_articles("712", "B1", on_progress=lambda m: None)
    check("매물 수집: 현대1,2차 전세 4페이지 64건", len(arts) == 64, f"{len(arts)}건")
    check("정확한 전용면적 부착 (196.21 등)", any(a.get("exclusiveArea") == 196.21 for a in arts))
    arts = nl.get_articles("104126", "A1:B1", on_progress=lambda m: None)
    n59 = len(c.filter_by_area_types(arts, ["59타입"]))
    n84 = len(c.filter_by_area_types(arts, ["84타입"]))
    check("24단지 매매+전세 493건 / 59타입 101 / 84타입 117", (len(arts), n59, n84) == (493, 101, 117),
          f"{len(arts)}/{n59}/{n84}")
    for x, want in [(57.99, False), (58.88, True), (59.01, True), (60.0, False)]:
        check(f"59타입 경계 {x}", bool(c.filter_by_area_types([{"exclusiveArea": x}], ["59타입"])) == want)
    for x, want in [(82.99, False), (84.09, True), (85.99, True), (86.0, False)]:
        check(f"84타입 경계 {x}", bool(c.filter_by_area_types([{"exclusiveArea": x}], ["84타입"])) == want)
    check("가격 변환 '6억 7,000' -> 67000", c.price_to_manwon("6억 7,000") == 67000)
    check("가격 변환 '15.5억' -> 155000", c.price_to_manwon("15.5억") == 155000)

    cart = c.Cart(os.path.join(OUT, "cart_backend.json"))
    cart.clear()
    a1 = nl.get_articles("108582", "A1:B1", on_progress=lambda m: None, real_estate_type="OPST:OBYG")
    a2 = nl.get_articles("111601", "A1:B1", on_progress=lambda m: None, real_estate_type="OPST:OBYG")
    cart.add("108582", "위례효성해링턴타워", "매매+전세 / 전체면적", a1)
    cart.add("111601", "효성해링턴타워더퍼스트", "매매+전세 / 전체면적", a2)
    added, dup = cart.add("111601", "효성해링턴타워더퍼스트", "매매+전세 / 전체면적", a2)
    check("장바구니: 191 + 220 = 411건, 재담기 중복 0건", cart.total() == 411 and added == 0 and dup == 220)
    check("장바구니: 재시작 후 유지", c.Cart(os.path.join(OUT, "cart_backend.json")).total() == 411)
    xl = cart.export_xlsx(os.path.join(OUT, "cart_backend.xlsx"))
    import openpyxl
    wb = openpyxl.load_workbook(xl)
    check("장바구니 엑셀: 요약/전체/단지별 시트, 전체 411행",
          wb.sheetnames[:2] == ["요약", "전체"] and wb["전체"].max_row - 1 == 411, str(wb.sheetnames))


def test_frontend():
    print("\n[프론트엔드]")
    from PySide6.QtCore import QTimer, Qt
    from PySide6.QtGui import QFont
    from PySide6.QtWidgets import QApplication, QMessageBox
    import naver_land_gui as g

    xl = os.path.join(OUT, "cart_gui.xlsx")
    g.QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (xl, ""))
    QMessageBox.exec = lambda self: 0
    QMessageBox.information = staticmethod(lambda *a, **k: None)
    app = QApplication.instance() or QApplication(sys.argv)
    app.setFont(QFont("Malgun Gothic", 9))
    app.setStyleSheet(g.TOGGLE_STYLE)
    w = g.MainWindow()
    w.resize(1400, 800)
    w.show()
    click = lambda s: w.tbl.horizontalHeader().sectionClicked.emit(s)  # noqa: E731

    def pick(no):
        idx = [i for i in range(w.tbl_cx.rowCount()) if w.tbl_cx.item(i, 4).text() == no][0]
        w.tbl_cx.selectRow(idx)
        w.fetch()

    def s1():
        w.keyword.setText("위례")
        w.search()

    def s2():
        check("종류 버튼: 둘 다 켜면 110개", w.tbl_cx.rowCount() == 110)
        w.type_btns["아파트"].setChecked(False)
        check("종류 버튼: 오피스텔만 24개", w.tbl_cx.rowCount() == 24)
        pick("108582")

    def s3():
        n = w.tbl.rowCount()
        check("위례효성해링턴타워 매매+전세 191건 표시", n == 191, str(n))
        check("처음엔 체크 0건, 담기 버튼 꺼짐", len(w.checked_articles()) == 0 and not w.btn_cart_add.isEnabled())
        click(0)
        check("전체선택 -> 191건 체크", len(w.checked_articles()) == 191)
        click(0)
        check("전체선택 다시 -> 0건", len(w.checked_articles()) == 0)
        for r in (0, 2, 4):
            w.tbl.item(r, 0).setCheckState(Qt.Checked)
        before = sorted(a["articleNo"] for a in w.checked_articles())
        click(w.col("가격(만원)"))
        after = sorted(a["articleNo"] for a in w.checked_articles())
        check("정렬해도 체크 유지", before == after and len(after) == 3)
        w.cart_add()
        check("체크한 3건만 담김 + '담김' 표시",
              w.cart.total() == 3 and sum(w.tbl.item(r, 0).text() == "담김" for r in range(n)) == 3)
        w.grab().save(os.path.join(OUT, "gui_checkbox.png"))
        w.keyword.setText("미사")
        w.search()

    def s4():
        pick("111601")

    def s5():
        w.area_btns["84타입"].setChecked(True)
        click(0)
        w.cart_add()
        w.cart_download()
        import openpyxl
        wb = openpyxl.load_workbook(xl)
        check("미사 84타입 2건 추가 -> 엑셀 전체 5행, 시트 4개",
              wb["전체"].max_row - 1 == 5 and len(wb.sheetnames) == 4, str([(s.title, s.max_row - 1) for s in wb]))
        app.quit()

    for ms, f in [(300, s1), (1500, s2), (3500, s3), (5000, s4), (7500, s5)]:
        QTimer.singleShot(ms, f)
    app.exec()


def test_safety():
    """v7: 조회 중에 다른 조회/검색/창 닫기를 해도 꺼지지 않는지 (v6까지는 0xC0000409로 강제 종료됨)"""
    print("\n[안정성 (v7)]")
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QApplication, QMessageBox
    import naver_land_gui as g

    app = QApplication.instance() or QApplication(sys.argv)
    warnings = []
    QMessageBox.warning = staticmethod(lambda *a, **k: warnings.append(a[2] if len(a) > 2 else ""))
    w = g.MainWindow()
    w.show()
    w.api._wait = lambda sec: _orig_wait(w.api, 0.2)   # 페이지마다 0.2초 -> 25페이지 조회가 몇 초 걸림
    st = {}

    def s1():
        w.all_complexes = w.api.search_complexes("위례")
        w.type_btns["오피스텔"].setChecked(False)
        rows = [i for i in range(w.tbl_cx.rowCount()) if w.tbl_cx.item(i, 4).text() == "104126"]
        w.tbl_cx.selectRow(rows[0])
        st["other"] = 0 if rows[0] else 1
        w.fetch()
        st["worker"] = w.worker
        check("조회 시작 -> 작업 중, [중지] 켜짐, 검색창/단지목록 잠김",
              w.is_busy() and w.btn_stop.isEnabled() and not w.keyword.isEnabled() and not w.tbl_cx.isEnabled())

    def s2():
        w.tbl_cx.cellDoubleClicked.emit(st["other"], 0)        # v6: 여기서 꺼짐
        check("조회 중 단지 더블클릭 -> 무시하고 안내 (꺼지지 않음)",
              w.worker is st["worker"] and "조회 중" in w.status.text())
        w.keyword.setText("미사")
        w.keyword.returnPressed.emit()                         # v6: 여기서도 꺼짐
        check("조회 중 검색창 Enter -> 무시하고 안내", w.worker is st["worker"] and w.is_busy())

    def s3():
        w.stop()

    def s4():
        n = w.tbl.rowCount()
        check("[중지] -> 받은 데까지만 표시 (0 < 표시 < 493건)", 0 < n < 493 and not w.is_busy(), f"{n}건")
        check("[중지] 후 안내 문구 + 버튼 다시 켜짐",
              w.status.text().startswith("중지했어요") and w.btn_fetch.isEnabled() and w.tbl_cx.isEnabled())

        orig_save, orig_dlg = g.save, g.QFileDialog.getSaveFileName
        g.QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (os.path.join(OUT, "x.xlsx"), ""))
        g.save = lambda *a, **k: (_ for _ in ()).throw(PermissionError("열려 있음"))
        w.export()
        g.save, g.QFileDialog.getSaveFileName = orig_save, orig_dlg
        check("현재 표 엑셀 저장: 파일이 열려 있으면 안내창", warnings and "열려 있어요" in warnings[-1])

        logp = os.path.join(OUT, "error.log")
        if os.path.exists(logp):
            os.remove(logp)
        g.setup_error_log(logp)
        try:
            raise ValueError("테스트 오류")
        except ValueError:
            sys.excepthook(*sys.exc_info())
        for h in list(g.log.handlers):
            h.flush()
        text = open(logp, encoding="utf-8").read()
        check("놓친 오류 -> error.log 기록 + 안내창",
              "ValueError: 테스트 오류" in text and "error.log" in warnings[-1])

        w.fetch()                                              # 다시 조회 시작 후 바로 창 닫기

    def s5():
        st["busy_at_close"] = w.is_busy()
        w.close()                                              # v6: 닫을 때 강제 종료
        st["hidden"] = w.isHidden()

    def on_quit():
        st["quit_running"] = w.worker.isRunning()
        st["quit_by_close"] = "fallback" not in st

    app.aboutToQuit.connect(on_quit)
    for ms, f in [(300, s1), (900, s2), (1500, s3), (2800, s4), (3400, s5)]:
        QTimer.singleShot(ms, f)
    QTimer.singleShot(20000, lambda: (st.__setitem__("fallback", True), app.quit()))
    app.exec()
    app.aboutToQuit.disconnect(on_quit)
    check("조회 중 창 닫기 -> 창은 바로 숨고, 작업이 멈춘 뒤 정상 종료",
          st.get("busy_at_close") and st.get("hidden") and st.get("quit_by_close") and st.get("quit_running") is False,
          str(st.get("quit_by_close")))


if __name__ == "__main__":
    test_backend()
    try:
        test_frontend()
        test_safety()
    except ImportError as e:
        print("  [SKIP] 프론트엔드 테스트 (PySide6 없음):", e)
    fails = [n for n, ok in RESULTS if not ok]
    print(f"\n결과: {len(RESULTS) - len(fails)}/{len(RESULTS)} 통과" + (f", 실패: {fails}" if fails else ""))
    if MISSING:
        print("fixture 없음:", MISSING[:3])
    sys.exit(1 if fails else 0)
