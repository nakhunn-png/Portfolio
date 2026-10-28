# -*- coding: utf-8 -*-
"""
네이버 부동산 아파트 매물 크롤러
- 영상(윤자동, 커서AI로 네이버 부동산 매물 자동 수집하기)의 '네트워크 요청 재현' 방식을 파이썬으로 구현
- 차이점: cURL을 복사해 토큰을 하드코딩하지 않고, 페이지 HTML에서 토큰(3시간 유효)을 자동으로 받아옴

필요 라이브러리: pip install requests openpyxl
사용법:
    python naver_land_crawler.py                 # 대화형 (단지 검색 -> 선택 -> 수집)
    python naver_land_crawler.py "압구정 현대"   # 검색어 바로 지정
    python naver_land_crawler.py --complex 3037 --trade A1   # 단지번호로 바로 수집 (A1=매매, B1=전세, B2=월세)
"""
import argparse
import base64
import csv
import json
import random
import re
import sys
import threading
import time
from datetime import datetime
from urllib.parse import quote

# 크롬 브라우저와 똑같은 방식(TLS 지문)으로 접속하는 curl_cffi 우선 사용, 없으면 requests
try:
    from curl_cffi import requests as http
    HTTP_LIB = "curl_cffi"
except ImportError:  # pragma: no cover
    import requests as http
    HTTP_LIB = "requests"

BASE = "https://new.land.naver.com"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/141.0.0.0 Safari/537.36")
TRADE_TYPES = {"": "전체", "A1": "매매", "B1": "전세", "B2": "월세"}
# 매물 종류 그룹 (네이버 코드)
PROPERTY_GROUPS = {
    "아파트": ["APT", "ABYG", "JGC"],      # 아파트, 아파트분양권, 재건축
    "오피스텔": ["OPST", "OBYG"],         # 오피스텔, 오피스텔분양권
}
REAL_ESTATE_TYPE = "APT:ABYG:JGC"
# 전용면적 타입: (이상, 미만) ㎡  -> 59타입은 58.88, 59.01 등, 84타입은 84.09, 84.99 등 포함
AREA_TYPES = {
    "59타입": (58.0, 60.0),
    "84타입": (83.0, 86.0),
}


def trade_label(code):
    """'A1:B1' -> '매매+전세', '' -> '전체'"""
    return "+".join(TRADE_TYPES.get(c, c) for c in code.split(":")) if code else "전체"


def exclusive_area(article):
    """정확한 전용면적(소수점). 없으면 목록의 정수값(area2) 사용"""
    v = article.get("exclusiveArea")
    if v is None:
        v = article.get("area2")
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def filter_by_area_types(articles, area_types):
    """area_types: ['59타입', '84타입'] 중 선택. 비어 있으면 전체 반환"""
    if not area_types:
        return list(articles)
    ranges = [AREA_TYPES[t] for t in area_types]
    out = []
    for a in articles:
        x = exclusive_area(a)
        if x is not None and any(lo <= x < hi for lo, hi in ranges):
            out.append(a)
    return out


def property_group(complex_info):
    """단지 정보 -> '아파트' / '오피스텔' / '기타'"""
    code = complex_info.get("realEstateTypeCode", "")
    for name, codes in PROPERTY_GROUPS.items():
        if code in codes:
            return name
    return "기타"


class Cancelled(Exception):
    """사용자가 [중지]를 누름. partial = 그때까지 받은 결과"""

    def __init__(self, partial=None):
        super().__init__("사용자가 중지함")
        self.partial = partial


class NaverLand:
    def __init__(self, delay=(1.5, 3.0), timeout=20):
        if HTTP_LIB == "curl_cffi":
            self.s = http.Session(impersonate="chrome")
        else:
            self.s = http.Session()
            self.s.headers["User-Agent"] = UA
        self.s.headers.update({
            "Accept-Language": "ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7",
        })
        self.delay = delay
        self.timeout = timeout
        self.token = None
        self.token_exp = 0
        self._cancel = threading.Event()

    # ---------- 중지 ----------
    def cancel(self):
        """다른 스레드(화면)에서 호출. 진행 중인 수집을 다음 대기 시점에 멈춤"""
        self._cancel.set()

    def reset_cancel(self):
        self._cancel.clear()

    def _wait(self, sec):
        """sec초 대기. 도중에 중지를 누르면 바로 Cancelled"""
        if self._cancel.wait(sec):
            raise Cancelled()

    def _raw_get(self, url, step, headers, tries=3):
        """시간 초과/연결 오류 시 자동 재시도. 실패하면 어느 단계인지 알려줌"""
        last = None
        for i in range(tries):
            try:
                return self.s.get(url, headers=headers, timeout=self.timeout)
            except Exception as e:  # noqa: BLE001  (timeout, connection error 등)
                last = e
                self._wait(3 * (i + 1))
        raise RuntimeError(
            f"[{step}] 네이버 서버가 응답하지 않습니다 ({tries}번 시도).\n"
            f"인터넷 연결을 확인하고, 크롬에서 new.land.naver.com 이 열리는지 확인해 주세요.\n"
            f"(접속 방식: {HTTP_LIB}, 상세: {type(last).__name__}: {last})")

    # ---------- 토큰 ----------
    def refresh_token(self):
        r = self._raw_get(BASE + "/complexes", "1단계: 네이버 부동산 접속/토큰 받기", {
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Sec-Fetch-Dest": "document", "Sec-Fetch-Mode": "navigate",
            "Sec-Fetch-Site": "none", "Upgrade-Insecure-Requests": "1",
        })
        if r.status_code != 200:
            raise RuntimeError(f"[1단계: 토큰 받기] 네이버가 접속을 거부했습니다 (HTTP {r.status_code}).")
        m = re.search(r'"token"\s*:\s*\{\s*"token"\s*:\s*"([\w-]+\.[\w-]+\.[\w-]+)"', r.text) \
            or re.search(r"eyJ[\w-]+\.[\w-]+\.[\w-]+", r.text)
        if not m:
            raise RuntimeError("[1단계: 토큰 받기] 토큰을 찾지 못했습니다. 네이버가 차단 페이지를 보냈거나 "
                               "페이지 구조가 바뀌었을 수 있어요.")
        self.token = m.group(1) if m.groups() else m.group(0)
        payload = self.token.split(".")[1]
        payload += "=" * (-len(payload) % 4)
        self.token_exp = json.loads(base64.urlsafe_b64decode(payload)).get("exp", time.time() + 3600)
        self.s.headers["authorization"] = "Bearer " + self.token

    def _ensure_token(self):
        if not self.token or time.time() > self.token_exp - 300:  # 만료 5분 전 갱신
            self.refresh_token()

    # ---------- 공통 GET (대기 + 재시도) ----------
    def _get(self, path, retries=4, step="데이터 요청"):
        self._ensure_token()
        for attempt in range(retries):
            self._wait(random.uniform(*self.delay))
            r = self._raw_get(BASE + path, step, {
                "Accept": "application/json, text/plain, */*",
                "Referer": BASE + "/complexes",
                "Sec-Fetch-Dest": "empty", "Sec-Fetch-Mode": "cors", "Sec-Fetch-Site": "same-origin",
            })
            if r.status_code == 200:
                return r.json()
            if r.status_code == 401:          # 토큰 만료
                self.refresh_token()
                continue
            if r.status_code == 429:          # 요청 과다 -> 점점 길게 쉬기
                wait = 15 * (attempt + 1)
                print(f"  [429] 요청이 많아 {wait}초 쉽니다...")
                self._wait(wait)
                continue
            raise RuntimeError(f"[{step}] 네이버가 요청을 거부했습니다 (HTTP {r.status_code}).")
        raise RuntimeError(f"[{step}] 여러 번 시도했지만 실패했습니다. 잠시 후 다시 해주세요.")

    # ---------- 단지 검색 ----------
    def search_complexes(self, keyword, max_pages=5, on_progress=lambda m: None):
        """검색 결과를 끝 페이지까지(최대 max_pages) 모두 가져옴. 페이지당 50개"""
        results = []
        try:
            for page in range(1, max_pages + 1):
                data = self._get(f"/api/search?keyword={quote(keyword)}&page={page}",
                                 step=f"2단계: 단지 검색 {page}페이지")
                results.extend(data.get("complexes", []))
                on_progress(f"단지 검색 중... {len(results)}개")
                if not data.get("isMoreData"):
                    break
        except Cancelled:
            raise Cancelled(results) from None
        return results

    # ---------- 단지의 평형별 정확한 전용면적 ----------
    def get_complex_areas(self, complex_no):
        """{'74C': 34.84, '84A': 84.09, ...}  (평형이름 -> 전용면적)"""
        data = self._get(f"/api/complexes/{complex_no}?sameAddressGroup=false",
                         step="3단계: 단지 면적 정보")
        areas = {}
        for p in data.get("complexPyeongDetailList") or []:
            try:
                areas[str(p.get("pyeongName"))] = float(p.get("exclusiveArea"))
            except (TypeError, ValueError):
                pass
        return areas

    def attach_exact_areas(self, complex_no, articles):
        """각 매물에 exclusiveArea(정확한 전용면적) 추가. 실패해도 수집은 계속"""
        try:
            areas = self.get_complex_areas(complex_no)
        except Cancelled:
            raise
        except Exception:  # noqa: BLE001
            areas = {}
        for a in articles:
            name = str(a.get("areaName", ""))
            if name in areas:
                a["exclusiveArea"] = areas[name]
        return articles

    # ---------- 매물 목록 (전체 페이지) ----------
    def get_articles(self, complex_no, trade_type="", max_pages=50, on_progress=print, real_estate_type=None):
        """trade_type: '' 전체, 'A1' 매매, 'A1:B1' 매매+전세 처럼 ':'로 여러 개 가능"""
        ret = real_estate_type or REAL_ESTATE_TYPE
        all_items, page = [], 1
        try:
            while page <= max_pages:
                path = (f"/api/articles/complex/{complex_no}?realEstateType={quote(ret)}"
                        f"&tradeType={quote(trade_type)}&tag=%3A%3A%3A%3A%3A%3A%3A%3A"
                        f"&rentPriceMin=0&rentPriceMax=900000000&priceMin=0&priceMax=900000000"
                        f"&areaMin=0&areaMax=900000000&showArticle=false&sameAddressGroup=false"
                        f"&priceType=RETAIL&directions=&page={page}&complexNo={complex_no}"
                        f"&buildingNos=&areaNos=&type=list&order=rank")
                data = self._get(path, step=f"3단계: 매물 {page}페이지 수집")
                items = data.get("articleList", [])
                all_items.extend(items)
                on_progress(f"  {page}페이지: {len(items)}건 (누적 {len(all_items)}건)")
                if not data.get("isMoreData"):
                    break
                page += 1
            return self.attach_exact_areas(complex_no, all_items)
        except Cancelled:  # 중지 -> 받은 데까지 돌려줌 (정확한 면적은 생략, 목록의 정수 면적 사용)
            raise Cancelled(all_items) from None


# ---------- 가격 문자열 -> 만원 단위 숫자 ----------
def price_to_manwon(text):
    """'30억' -> 300000, '6억 7,000' -> 67000, '15.5억' -> 155000, '9,500' -> 9500"""
    if not text:
        return None
    t = str(text).replace(",", "").replace(" ", "")
    total = 0.0
    if "억" in t:
        eok, _, rest = t.partition("억")
        total += float(eok or 0) * 10000
        t = rest
    if t:
        try:
            total += float(t)
        except ValueError:
            return None
    return int(total)


COLUMNS = [
    ("매물번호", "articleNo"), ("단지명", "articleName"), ("동", "buildingName"),
    ("거래유형", "tradeTypeName"), ("매물종류", "realEstateTypeName"), ("가격(표시)", "dealOrWarrantPrc"), ("가격(만원)", None),
    ("월세(만원)", "rentPrc"), ("공급면적㎡", "area1"), ("전용면적㎡", None), ("평형명", "areaName"),
    ("평형(공급)", None), ("층", "floorInfo"), ("방향", "direction"),
    ("확인일자", "articleConfirmYmd"), ("특징", "articleFeatureDesc"), ("태그", "tagList"),
    ("중개사", "realtorName"), ("정보제공", "cpName"), ("링크", None),
]


def to_rows(articles, complex_no=""):
    rows = []
    for a in articles:
        row = []
        for name, key in COLUMNS:
            if name == "가격(만원)":
                v = price_to_manwon(a.get("dealOrWarrantPrc"))
            elif name == "전용면적㎡":
                v = exclusive_area(a)
            elif name == "평형(공급)":
                v = round(a["area1"] / 3.3058, 1) if a.get("area1") else None
            elif name == "링크":
                v = f"{BASE}/complexes/{a.get('_complexNo') or complex_no}?articleNo={a.get('articleNo')}"
            elif key == "tagList":
                v = ", ".join(a.get("tagList") or [])
            else:
                v = a.get(key)
            row.append(v)
        rows.append(row)
    return rows


COL_WIDTHS = [12, 18, 8, 8, 9, 12, 11, 9, 10, 10, 9, 9, 8, 7, 10, 40, 30, 24, 12, 45]


def _write_sheet(ws, headers, rows, widths):
    from openpyxl.utils import get_column_letter
    ws.append(headers)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def save(articles, basename, complex_no="", write_json=True):
    headers = [c[0] for c in COLUMNS]
    rows = to_rows(articles, complex_no)
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "매물"
        _write_sheet(ws, headers, rows, COL_WIDTHS)
        path = basename + ".xlsx"
        wb.save(path)
    except ImportError:
        path = basename + ".csv"
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(headers)
            w.writerows(rows)
    if write_json:
        with open(basename + ".json", "w", encoding="utf-8") as f:
            json.dump(articles, f, ensure_ascii=False, indent=2)
    return path


# ---------- 엑셀 장바구니 ----------
class Cart:
    """여러 단지의 조회 결과를 담아두었다가 하나의 엑셀로 저장.
    - 같은 매물(매물번호)은 한 번만 담김
    - path를 주면 담을 때마다 자동 저장 -> 프로그램을 꺼도 유지
    """

    def __init__(self, path=None):
        self.path = path
        self.groups = []  # [{complexNo, complexName, condition, addedAt, articles}]
        if path:
            self.load()

    # 저장/불러오기
    def load(self):
        try:
            with open(self.path, encoding="utf-8") as f:
                self.groups = json.load(f).get("groups", [])
        except (OSError, ValueError):
            self.groups = []

    def persist(self):
        if not self.path:
            return
        try:
            import os
            os.makedirs(os.path.dirname(self.path) or ".", exist_ok=True)
            with open(self.path, "w", encoding="utf-8") as f:
                json.dump({"groups": self.groups}, f, ensure_ascii=False)
        except OSError:
            pass

    # 조작
    def article_ids(self):
        return {a.get("articleNo") for g in self.groups for a in g["articles"]}

    def add(self, complex_no, complex_name, condition, articles):
        """반환: (새로 담은 수, 중복라 건너뛴 수)"""
        have = self.article_ids()
        new = []
        for a in articles:
            if a.get("articleNo") in have:
                continue
            b = dict(a)
            b["_complexNo"] = complex_no
            new.append(b)
            have.add(a.get("articleNo"))
        if new:
            for g in self.groups:  # 같은 단지+같은 조건이면 합치기
                if g["complexNo"] == complex_no and g["condition"] == condition:
                    g["articles"].extend(new)
                    break
            else:
                self.groups.append({
                    "complexNo": complex_no, "complexName": complex_name, "condition": condition,
                    "addedAt": datetime.now().strftime("%Y-%m-%d %H:%M"), "articles": new,
                })
            self.persist()
        return len(new), len(articles) - len(new)

    def remove(self, index):
        del self.groups[index]
        self.persist()

    def clear(self):
        self.groups = []
        self.persist()

    def total(self):
        return sum(len(g["articles"]) for g in self.groups)

    def summary(self):
        return f"장바구니: {self.total()}건 (단지 {len({g['complexNo'] for g in self.groups})}곳)"

    # 엑셀 내보내기: [요약] + [전체] + 담은 묶음별 시트
    def export_xlsx(self, path):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        headers = [c[0] for c in COLUMNS]
        wb = Workbook()

        ws = wb.active
        ws.title = "요약"
        _write_sheet(ws, ["번호", "단지명", "단지번호", "담은 조건", "매물 수", "담은 시각", "시트"],
                     [[i, g["complexName"], g["complexNo"], g["condition"], len(g["articles"]),
                       g["addedAt"], self._sheet_name(i, g)] for i, g in enumerate(self.groups, 1)]
                     + [["", "합계", "", "", self.total(), "", ""]],
                     [6, 24, 10, 26, 9, 16, 24])
        ws.cell(row=len(self.groups) + 2, column=2).font = Font(bold=True)

        ws = wb.create_sheet("전체")
        rows = []
        for i, g in enumerate(self.groups, 1):
            for r in to_rows(g["articles"], g["complexNo"]):
                rows.append([i, g["condition"]] + r)
        _write_sheet(ws, ["번호", "담은 조건"] + headers, rows, [6, 20] + COL_WIDTHS)

        for i, g in enumerate(self.groups, 1):
            ws = wb.create_sheet(self._sheet_name(i, g))
            _write_sheet(ws, headers, to_rows(g["articles"], g["complexNo"]), COL_WIDTHS)
        wb.save(path)
        return path

    @staticmethod
    def _sheet_name(i, g):
        name = re.sub(r'[\\/*?:\[\]]', "_", f"{i}_{g['complexName']}")
        return name[:31]


def main():
    ap = argparse.ArgumentParser(description="네이버 부동산 아파트 매물 크롤러")
    ap.add_argument("keyword", nargs="?", help="단지/지역 검색어 (예: 압구정 현대)")
    ap.add_argument("--complex", help="단지번호 (검색 생략)")
    ap.add_argument("--trade", default=None, choices=list(TRADE_TYPES), help="A1=매매 B1=전세 B2=월세, 생략=전체")
    ap.add_argument("--area", nargs="*", default=None, choices=list(AREA_TYPES),
                    help="전용면적 타입 (예: --area 59타입 84타입), 생략=전체")
    args = ap.parse_args()

    nl = NaverLand()
    complex_no, complex_name = args.complex, args.complex

    if not complex_no:
        keyword = args.keyword or input("검색어 (예: 압구정 현대): ").strip()
        complexes = nl.search_complexes(keyword)
        if not complexes:
            print("검색 결과가 없습니다.")
            return
        for i, c in enumerate(complexes, 1):
            print(f"{i:2}. {c['complexName']} ({c.get('cortarAddress','')}, "
                  f"{c.get('totalHouseholdCount','?')}세대, {c.get('realEstateTypeName','')}) #{c['complexNo']}")
        sel = int(input("번호 선택: ").strip() or 1)
        complex_no = complexes[sel - 1]["complexNo"]
        complex_name = complexes[sel - 1]["complexName"]

    trade = args.trade
    if trade is None:
        t = input("거래유형 (Enter=전체, 1=매매, 2=전세, 3=월세): ").strip()
        trade = {"1": "A1", "2": "B1", "3": "B2"}.get(t, "")

    print(f"\n[{complex_name}] {TRADE_TYPES[trade]} 매물 수집 시작...")
    articles = nl.get_articles(complex_no, trade)
    area = args.area
    if area is None:
        t = input("전용면적 (Enter=전체, 1=59타입, 2=84타입, 12=둘 다): ").strip()
        area = [n for k, n in (("1", "59타입"), ("2", "84타입")) if k in t]
    if area:
        before = len(articles)
        articles = filter_by_area_types(articles, area)
        print(f"  면적 필터({'+'.join(area)}): {before}건 -> {len(articles)}건")
    if not articles:
        print("매물이 없습니다.")
        return
    safe = re.sub(r'[\\/:*?"<>|]', "_", str(complex_name))
    path = save(articles, f"매물_{safe}_{TRADE_TYPES[trade]}_{datetime.now():%Y%m%d_%H%M}", complex_no)
    print(f"\n완료! {len(articles)}건 저장 -> {path}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(0)
